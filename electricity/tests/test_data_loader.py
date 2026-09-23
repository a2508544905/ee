"""data_loader 与 report_exporter 模块单元测试"""

import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from lib.tariff_manager import TariffManager  # noqa: E402
from lib.calculator import Calculator  # noqa: E402
from lib.data_loader import DataLoader  # noqa: E402
from lib import report_exporter  # noqa: E402


class TestDataLoader(unittest.TestCase):
    """测试 DataLoader 的 CSV 导入、单条录入与 JSON 持久化"""

    def setUp(self):
        self.dl = DataLoader(Calculator(TariffManager()))
        self.tmp = tempfile.mkdtemp()

    def teardown_cleanup(self):
        return

    def _write_csv(self, text):
        """在临时目录写入 CSV 并返回路径。"""
        path = os.path.join(self.tmp, "sample.csv")
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)
        return path

    def test_批量导入逐条算费(self):
        path = self._write_csv("用户名,月份,地区,用电量\nA,1,贵州,120\nB,2,贵州,300\n")
        recs = self.dl.load_csv(path, default_region="贵州")
        self.assertEqual(len(recs), 2)
        self.assertEqual(recs[0]["username"], "A")
        self.assertAlmostEqual(recs[0]["total"], 54.67, places=2)  # 120 @0.4556
        self.assertEqual(recs[1]["username"], "B")

    def test_单条录入自动计算电费(self):
        rec = self.dl.add_record("C", 3, "贵州", 200)
        self.assertEqual(rec["username"], "C")
        self.assertEqual(rec["month"], 3)
        self.assertAlmostEqual(rec["total"], 91.12, places=2)  # 200 @0.4556
        self.assertIn("tiers", rec)

    def test_单条录入负数报错(self):
        with self.assertRaises(ValueError):
            self.dl.add_record("D", 1, "贵州", -1)

    def test_json持久化往返(self):
        path = os.path.join(self.tmp, "records.json")
        dl = DataLoader(Calculator(TariffManager()), records_path=path)
        recs = dl.load_csv(self._write_csv("用户名,月份,地区,用电量\nA,1,贵州,120\n"), "贵州")
        dl.save_records(recs)
        loaded = dl.load_records()
        self.assertEqual(len(loaded), 1)
        self.assertEqual(loaded[0]["username"], "A")

    def test_json损坏容错(self):
        path = os.path.join(self.tmp, "records.json")
        with open(path, "w", encoding="utf-8") as f:
            f.write("{ 损坏的 json")
        dl = DataLoader(Calculator(TariffManager()), records_path=path)
        self.assertEqual(dl.load_records(), [])

    def test_json文件不存在返回空(self):
        path = os.path.join(self.tmp, "none.json")
        dl = DataLoader(Calculator(TariffManager()), records_path=path)
        self.assertEqual(dl.load_records(), [])


class TestReportExporter(unittest.TestCase):
    """测试 CSV / Excel 导出"""

    def setUp(self):
        self.dl = DataLoader(Calculator(TariffManager()))
        self.tmp = tempfile.mkdtemp()

    def _records(self):
        path = os.path.join(self.tmp, "s.csv")
        with open(path, "w", encoding="utf-8") as f:
            f.write("用户名,月份,地区,用电量\nA,1,贵州,120\nA,2,贵州,300\nB,3,贵州,500\n")
        return self.dl.load_csv(path, default_region="贵州")

    def test_导出csv(self):
        path = os.path.join(self.tmp, "out.csv")
        n = report_exporter.export_csv(self._records(), path)
        with open(path, "r", encoding="utf-8-sig") as f:
            content = f.read()
        self.assertIn("用户名", content)
        self.assertEqual(n, 3)

    def test_导出excel多sheet(self):
        path = os.path.join(self.tmp, "out.xlsx")
        report_exporter.export_excel(self._records(), path)
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("未安装 openpyxl")
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["原始数据", "用户汇总"])

    def test_按扩展名导出(self):
        p = os.path.join(self.tmp, "a.csv")
        n = report_exporter.export(self._records(), p)
        self.assertEqual(n, 3)


class TestLoaderErrorHandling(unittest.TestCase):
    """测试 CSV 导入的错误处理：跳过非法行 + 汇总报告 + 编码检测"""

    def setUp(self):
        self.dl = DataLoader(Calculator(TariffManager()))
        self.tmp = tempfile.mkdtemp()

    def _write_csv(self, text, encoding="utf-8"):
        path = os.path.join(self.tmp, "t.csv")
        with open(path, "w", encoding=encoding, newline="") as f:
            f.write(text)
        return path

    def test_跳过无效行并汇总(self):
        """负数、非数字、超大用量、非法月份的行应被跳过，有效行正常导入。"""
        path = self._write_csv(
            "用户名,月份,地区,用电量\nA,1,贵州,120\nB,2,贵州,-5\n"
            "C,3,贵州,abc\nD,13,贵州,100\nE,4,贵州,999999\nF,5,贵州,300\n"
        )
        recs = self.dl.load_csv(path, "贵州")
        self.assertEqual(len(recs), 2)  # 只保留 A 和 F
        report = self.dl.last_report
        self.assertEqual(report["loaded"], 2)
        self.assertEqual(report["skipped"], 4)
        self.assertEqual(len(report["errors"]), 4)

    def test_文件不存在报错(self):
        with self.assertRaises(ValueError) as ctx:
            self.dl.load_csv(os.path.join(self.tmp, "none.csv"), "贵州")
        self.assertIn("不存在", str(ctx.exception))

    def test_GBK编码自动识别(self):
        """GBK 编码的 CSV 应能被自动识别并正常导入。"""
        # GBK 编码写入中文表头
        path = os.path.join(self.tmp, "gbk.csv")
        content = "用户名,月份,地区,用电量\nA,1,贵州,120\n"
        with open(path, "w", encoding="gbk", newline="") as f:
            f.write(content)
        recs = self.dl.load_csv(path, "贵州")
        self.assertEqual(len(recs), 1)
        self.assertEqual(recs[0]["username"], "A")
        self.assertAlmostEqual(recs[0]["total"], 54.67, places=2)

    def test_重复记录去重(self):
        """同一用户同一月份出现两次，只保留第一条并记录重复。"""
        path = self._write_csv(
            "用户名,月份,地区,用电量\nA,1,贵州,120\nA,1,贵州,300\n"
        )
        recs = self.dl.load_csv(path, "贵州")
        self.assertEqual(len(recs), 1)
        self.assertEqual(recs[0]["usage"], 120)
        self.assertEqual(len(self.dl.last_report["duplicates"]), 1)